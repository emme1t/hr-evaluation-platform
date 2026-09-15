import hashlib
import json
import posixpath
import re
import zlib
from io import BytesIO
from pathlib import PurePosixPath
from threading import Lock
from uuid import uuid4
from zipfile import BadZipFile, ZipFile

from django.db import IntegrityError, OperationalError, transaction
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from openpyxl.utils.exceptions import InvalidFileException
from xml.etree import ElementTree

from apps.evaluations.models import EvaluationProject

from ..models import SummaryWorkbookTemplate
from ..storage import PrivateStorageError, delete_private, read_private, write_private
from .common import (
    OutputBoundaryError,
    ReportingSnapshotUnavailable,
    frozen_project_name,
    safe_cell_text,
)


MAX_TEMPLATE_BYTES = 10 * 1024 * 1024
MAX_ZIP_MEMBERS = 256
MAX_ZIP_MEMBER_BYTES = 10 * 1024 * 1024
MAX_ZIP_TOTAL_BYTES = 50 * 1024 * 1024
MAX_COMPRESSION_RATIO = 100
MAX_WORKSHEET_ROWS = 10_000
MAX_WORKSHEET_COLUMNS = 64
MAX_WORKSHEET_CELLS = 100_000
MAX_DEFINED_NAMES = 128
MAX_MERGES = 1_000
MAX_TABLES = 32
MAX_FORMULAS = 10_000
MAX_CELL_TEXT = 32_767

APPROVED_SHEETS = ("Sheet1", "Sheet2", "Sheet3")
APPROVED_HEADERS = {
    "A1": "员工编号",
    "B1": "被评价人",
    "C1": "邮箱",
    "D1": "部门",
    "E1": "上级",
    "I1": "上级均分",
    "J1": "同部门均分",
    "K1": "跨部门均分",
    "L1": "总分",
    "M1": "状态",
}
APPROVED_TARGET_COLUMNS = ("A", "B", "C", "D", "E", "I", "J", "K", "L", "M")
UNSAFE_FORMULA_SYNTAX = re.compile(
    r"(?:\[[^\]]+\][^!]*!|\bDDE\s*\(|\|)", re.IGNORECASE
)
FORMULA_FUNCTION = re.compile(
    r"(?<![A-Z0-9_.])([A-Z_][A-Z0-9_.]*)\s*\(", re.IGNORECASE
)
SAFE_FORMULA_FUNCTIONS = frozenset({"SUM"})
XLM_NAME = re.compile(r"(?:AUTO[_ .]?(?:OPEN|CLOSE)|EXEC|CALL|REGISTER|XLM)", re.I)
_REGISTRATION_LOCK = Lock()


class SummaryTemplateError(ValueError):
    def __init__(self, message, code="SUMMARY_TEMPLATE_INVALID"):
        super().__init__(message)
        self.code = code


class TemplateStructureError(SummaryTemplateError):
    pass


def require_hr_admin(actor):
    if not actor or not getattr(actor, "is_authenticated", False):
        raise SummaryTemplateError("需要 HR 管理员权限", "HR_ADMIN_REQUIRED")
    if not actor.groups.filter(name="HR_ADMIN").exists():
        raise SummaryTemplateError("需要 HR 管理员权限", "HR_ADMIN_REQUIRED")


def _sha256(payload):
    return hashlib.sha256(payload).hexdigest()


def _normalized_member_name(name):
    if not isinstance(name, str) or not name or "\\" in name or "\x00" in name:
        raise SummaryTemplateError("XLSX 成员路径无效")
    pure = PurePosixPath(name)
    if pure.is_absolute() or any(part in ("", ".", "..") for part in pure.parts):
        raise SummaryTemplateError("XLSX 成员路径越界")
    normalized = posixpath.normpath(name)
    if normalized.startswith("../") or normalized == "..":
        raise SummaryTemplateError("XLSX 成员路径越界")
    return normalized


def _relationship_source_directory(member_name):
    if member_name == "_rels/.rels":
        return ""
    marker = "/_rels/"
    if marker not in member_name or not member_name.endswith(".rels"):
        return posixpath.dirname(member_name)
    prefix, relation_file = member_name.split(marker, 1)
    source_name = relation_file[:-5]
    return posixpath.dirname(posixpath.join(prefix, source_name))


def _validate_relationships(member_name, payload, member_names):
    try:
        root = ElementTree.fromstring(payload)
    except ElementTree.ParseError as exc:
        raise SummaryTemplateError("XLSX 关系文件格式错误") from exc
    source_directory = _relationship_source_directory(member_name)
    for relation in root:
        target = relation.attrib.get("Target")
        relation_type = relation.attrib.get("Type")
        if not target or not relation_type:
            raise SummaryTemplateError("XLSX 关系定义不完整")
        if relation.attrib.get("TargetMode", "").lower() == "external":
            raise SummaryTemplateError("XLSX 不允许外部链接")
        if re.match(r"^[a-z][a-z0-9+.-]*:", target, re.I) or target.startswith("//"):
            raise SummaryTemplateError("XLSX 不允许外部链接")
        resolved = posixpath.normpath(
            target.lstrip("/")
            if target.startswith("/")
            else posixpath.join(source_directory, target)
        )
        if resolved.startswith("../") or resolved == "..":
            raise SummaryTemplateError("XLSX 关系路径越界")
        if resolved not in member_names:
            raise SummaryTemplateError("XLSX 关系目标不存在")


def validate_xlsx_container(payload):
    if len(payload) > MAX_TEMPLATE_BYTES:
        raise SummaryTemplateError("汇总模板不能超过 10MB")
    try:
        archive = ZipFile(BytesIO(payload))
    except (BadZipFile, OSError) as exc:
        raise SummaryTemplateError("文件不是有效的 XLSX 容器") from exc
    infos = archive.infolist()
    if len(infos) > MAX_ZIP_MEMBERS:
        raise SummaryTemplateError("XLSX 成员数量超过资源边界")
    names = set()
    folded = set()
    total_size = 0
    for info in infos:
        name = _normalized_member_name(info.filename)
        if name in names or name.casefold() in folded:
            raise SummaryTemplateError("XLSX 包含重复成员")
        names.add(name)
        folded.add(name.casefold())
        if info.flag_bits & 0x1:
            raise SummaryTemplateError("XLSX 不允许加密成员")
        if info.file_size > MAX_ZIP_MEMBER_BYTES:
            raise SummaryTemplateError("XLSX 单个成员超过资源边界")
        total_size += info.file_size
        if total_size > MAX_ZIP_TOTAL_BYTES:
            raise SummaryTemplateError("XLSX 解压数据超过资源边界")
        if info.file_size and info.compress_size == 0:
            raise SummaryTemplateError("XLSX 压缩比超过资源边界")
        if info.compress_size and info.file_size / info.compress_size > MAX_COMPRESSION_RATIO:
            raise SummaryTemplateError("XLSX 压缩比超过资源边界")
        lowered = name.casefold()
        if "externallinks/" in lowered:
            raise SummaryTemplateError("XLSX 不允许外部链接")
        if lowered.endswith("vbaproject.bin") or "macrosheets/" in lowered:
            raise SummaryTemplateError("XLSX 不允许宏或 VBA")
    required = {"[Content_Types].xml", "_rels/.rels", "xl/workbook.xml"}
    if not required.issubset(names):
        raise SummaryTemplateError("XLSX 容器缺少必需成员")
    member_payloads = {}
    for info in infos:
        name = _normalized_member_name(info.filename)
        try:
            member_payload = archive.read(info)
        except (BadZipFile, RuntimeError, OSError, EOFError, zlib.error) as exc:
            raise SummaryTemplateError("XLSX 成员无法安全读取") from exc
        if len(member_payload) != info.file_size:
            raise SummaryTemplateError("XLSX 成员长度不一致")
        member_payloads[name] = member_payload
    content_types = member_payloads["[Content_Types].xml"].lower()
    if b"macroenabled" in content_types or b"vba" in content_types:
        raise SummaryTemplateError("XLSX 不允许宏或 VBA 内容类型")
    for info in infos:
        name = _normalized_member_name(info.filename)
        if name.endswith((".xml", ".rels")):
            member_payload = member_payloads[name]
            if name.endswith(".rels"):
                _validate_relationships(name, member_payload, names)
            else:
                try:
                    root = ElementTree.fromstring(member_payload)
                except ElementTree.ParseError as exc:
                    raise SummaryTemplateError("XLSX XML 格式错误") from exc
                if name.startswith("xl/worksheets/"):
                    _validate_declared_worksheet_dimension(root)
    return archive


def _validate_declared_worksheet_dimension(root):
    dimension = root.find("{*}dimension")
    if dimension is None or not dimension.attrib.get("ref"):
        return
    try:
        _min_col, _min_row, max_col, max_row = range_boundaries(
            dimension.attrib["ref"]
        )
    except (TypeError, ValueError) as exc:
        raise SummaryTemplateError("汇总模板工作表维度无效") from exc
    if max_row > MAX_WORKSHEET_ROWS or max_col > MAX_WORKSHEET_COLUMNS:
        raise SummaryTemplateError("汇总模板稀疏维度超过资源边界")


def _formula_text(cell):
    if cell.data_type != "f":
        return None
    value = cell.value
    text = value if isinstance(value, str) else getattr(value, "text", None)
    if not isinstance(text, str) or not text.strip():
        raise SummaryTemplateError("汇总模板公式格式不受支持")
    return text


def _validate_formula(formula):
    if not formula:
        return
    text = str(formula).lstrip()
    functions = {match.group(1).upper() for match in FORMULA_FUNCTION.finditer(text)}
    if UNSAFE_FORMULA_SYNTAX.search(text) or not functions.issubset(
        SAFE_FORMULA_FUNCTIONS
    ):
        raise SummaryTemplateError("汇总模板包含不安全公式")


def _semantic_xml(value):
    return ElementTree.tostring(value.to_tree(), encoding="unicode")


def _defined_name_manifest(workbook):
    names = list(workbook.defined_names.values())
    if len(names) > MAX_DEFINED_NAMES:
        raise SummaryTemplateError("汇总模板定义名称超过资源边界")
    result = []
    for item in names:
        name = item.name or ""
        text = item.attr_text or ""
        if XLM_NAME.search(name) or XLM_NAME.search(text):
            raise SummaryTemplateError("汇总模板包含 XLM 定义名称")
        _validate_formula(text)
        result.append(
            {
                "name": name,
                "attr_text": text,
                "local_sheet_id": item.localSheetId,
                "hidden": bool(item.hidden),
                "definition": _semantic_xml(item),
            }
        )
    return sorted(result, key=lambda row: (row["name"], row["local_sheet_id"] or -1))


def _table_manifest(worksheet):
    tables = list(worksheet.tables.values())
    if len(tables) > MAX_TABLES:
        raise SummaryTemplateError("汇总模板数据表超过资源边界")
    result = []
    for table in tables:
        columns = []
        for column in table.tableColumns:
            calculated = getattr(column.calculatedColumnFormula, "text", None)
            totals = getattr(column.totalsRowFormula, "text", None)
            _validate_formula(calculated)
            _validate_formula(totals)
            columns.append(
                {
                    "id": column.id,
                    "name": column.name,
                    "calculated": calculated,
                    "totals": totals,
                }
            )
        result.append(
            {
                "name": table.name,
                "display_name": table.displayName,
                "ref": table.ref,
                "columns": columns,
                "definition": _semantic_xml(table),
            }
        )
    return sorted(result, key=lambda row: row["name"])


def workbook_structure_manifest(workbook):
    if tuple(workbook.sheetnames) != APPROVED_SHEETS:
        raise SummaryTemplateError("汇总模板工作表不符合批准结构")
    summary = workbook["Sheet1"]
    for coordinate, expected in APPROVED_HEADERS.items():
        if summary[coordinate].value != expected:
            raise SummaryTemplateError("汇总模板表头不符合批准结构")

    employee_rows = {}
    for row in range(2, summary.max_row + 1):
        value = summary[f"A{row}"].value
        if value in (None, ""):
            continue
        key = str(value)
        if key in employee_rows:
            raise SummaryTemplateError("汇总模板员工编号行重复")
        employee_rows[key] = row

    defined_names = _defined_name_manifest(workbook)
    sheets = []
    total_cells = 0
    formula_count = 0
    merge_count = 0
    for worksheet in workbook.worksheets:
        if worksheet.max_row > MAX_WORKSHEET_ROWS or worksheet.max_column > MAX_WORKSHEET_COLUMNS:
            raise SummaryTemplateError("汇总模板稀疏维度超过资源边界")
        cells = []
        target_coordinates = {
            f"{column}{row}"
            for row in employee_rows.values()
            for column in APPROVED_TARGET_COLUMNS
        } if worksheet.title == "Sheet1" else set()
        for cell in sorted(worksheet._cells.values(), key=lambda value: (value.row, value.column)):
            total_cells += 1
            value = cell.value
            if isinstance(value, str) and len(value) > MAX_CELL_TEXT:
                raise SummaryTemplateError("单元格文本超过 Excel 32767 字符边界")
            formula = _formula_text(cell)
            if formula:
                formula_count += 1
                _validate_formula(formula)
                if cell.coordinate in target_coordinates:
                    raise SummaryTemplateError("汇总模板目标单元格不能包含公式")
            cells.append(
                {
                    "coordinate": cell.coordinate,
                    "style_id": cell.style_id,
                    "style": {
                        "font": _semantic_xml(cell.font),
                        "fill": _semantic_xml(cell.fill),
                        "border": _semantic_xml(cell.border),
                        "alignment": _semantic_xml(cell.alignment),
                        "protection": _semantic_xml(cell.protection),
                        "number_format": cell.number_format,
                    },
                    "data_type": (
                        None
                        if cell.coordinate in target_coordinates
                        else cell.data_type
                    ),
                    "formula": formula,
                    "value": (
                        None
                        if cell.coordinate in target_coordinates or formula
                        else value
                    ),
                }
            )
        if total_cells > MAX_WORKSHEET_CELLS or formula_count > MAX_FORMULAS:
            raise SummaryTemplateError("汇总模板单元格或公式超过资源边界")
        merges = sorted(str(item) for item in worksheet.merged_cells.ranges)
        merge_count += len(merges)
        if merge_count > MAX_MERGES:
            raise SummaryTemplateError("汇总模板合并单元格超过资源边界")
        columns = {
            key: {
                "width": dimension.width,
                "hidden": bool(dimension.hidden),
                "outline_level": dimension.outlineLevel,
            }
            for key, dimension in sorted(worksheet.column_dimensions.items())
        }
        rows = {
            str(key): {
                "height": dimension.height,
                "hidden": bool(dimension.hidden),
                "outline_level": dimension.outlineLevel,
            }
            for key, dimension in sorted(worksheet.row_dimensions.items())
        }
        sheets.append(
            {
                "title": worksheet.title,
                "max_row": worksheet.max_row,
                "max_column": worksheet.max_column,
                "dimension": worksheet.calculate_dimension(),
                "freeze_panes": str(worksheet.freeze_panes or ""),
                "merged_cells": merges,
                "column_dimensions": columns,
                "row_dimensions": rows,
                "tables": _table_manifest(worksheet),
                "cells": cells,
            }
        )
    return {
        "format": 1,
        "sheets": sheets,
        "defined_names": defined_names,
        "target_rows": employee_rows,
        "target_columns": list(APPROVED_TARGET_COLUMNS),
    }


def _load_validated_workbook(payload):
    validate_xlsx_container(payload)
    try:
        workbook = load_workbook(BytesIO(payload), data_only=False, keep_links=False)
    except (BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as exc:
        raise SummaryTemplateError("无法解析汇总模板") from exc
    manifest = workbook_structure_manifest(workbook)
    signature = _sha256(
        json.dumps(
            manifest,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
            default=str,
        ).encode("utf-8")
    )
    return workbook, manifest, signature


def register_summary_template(file_bytes, filename, actor):
    require_hr_admin(actor)
    if not isinstance(file_bytes, bytes):
        raise SummaryTemplateError("汇总模板必须以字节形式上传")
    if not isinstance(filename, str) or not filename.lower().endswith(".xlsx"):
        raise SummaryTemplateError("只接受 .xlsx 汇总模板")
    base_name = PurePosixPath(filename.replace("\\", "/")).name
    if base_name != filename or len(base_name) > 180:
        raise SummaryTemplateError("汇总模板文件名无效")
    _workbook, manifest, signature = _load_validated_workbook(file_bytes)
    digest = _sha256(file_bytes)
    relative_path = f"templates/{uuid4().hex}.xlsx"
    stored = False
    try:
        with _REGISTRATION_LOCK:
            with transaction.atomic():
                latest = (
                    SummaryWorkbookTemplate.objects.select_for_update()
                    .filter(filename=base_name)
                    .order_by("-version")
                    .first()
                )
                version = 1 if latest is None else latest.version + 1
                write_private(relative_path, file_bytes)
                stored = True
                return SummaryWorkbookTemplate.objects.create(
                    filename=base_name,
                    version=version,
                    storage_path=relative_path,
                    sha256=digest,
                    structure_signature=signature,
                    structure_manifest=manifest,
                    created_by=actor,
                )
    except PrivateStorageError as exc:
        raise SummaryTemplateError("汇总模板私有存储失败") from exc
    except (IntegrityError, OperationalError) as exc:
        if stored:
            delete_private(relative_path)
        raise SummaryTemplateError(
            "汇总模板版本登记冲突，请重试",
            "SUMMARY_TEMPLATE_REGISTRATION_CONFLICT",
        ) from exc
    except Exception:
        if stored:
            delete_private(relative_path)
        raise


def template_snapshot(template):
    return {
        "public_id": str(template.public_id),
        "version": template.version,
        "sha256": template.sha256,
        "structure_signature": template.structure_signature,
    }


def assign_summary_template(project, template, actor):
    require_hr_admin(actor)
    with transaction.atomic():
        locked_project = EvaluationProject.objects.select_for_update().get(pk=project.pk)
        locked_template = SummaryWorkbookTemplate.objects.get(pk=template.pk)
        if locked_project.status != EvaluationProject.Status.DRAFT:
            raise SummaryTemplateError("只有草稿项目可以绑定汇总模板")
        if not locked_template.is_active:
            raise SummaryTemplateError("汇总模板版本未启用")
        locked_project.summary_template = locked_template
        locked_project.summary_template_snapshot = template_snapshot(locked_template)
        locked_project.save(
            update_fields=["summary_template", "summary_template_snapshot"]
        )
        project.summary_template_id = locked_template.pk
        project.summary_template_snapshot = locked_project.summary_template_snapshot
        return locked_project


def read_template_bytes(template):
    try:
        payload = read_private(template.storage_path)
    except PrivateStorageError as exc:
        raise TemplateStructureError("冻结汇总模板文件不可用") from exc
    if _sha256(payload) != template.sha256:
        raise TemplateStructureError("冻结汇总模板哈希不一致")
    return payload


def verify_project_template_binding(project):
    if not project.summary_template_id:
        raise SummaryTemplateError("项目必须先绑定启用的汇总模板")
    template = SummaryWorkbookTemplate.objects.get(pk=project.summary_template_id)
    expected_snapshot = template_snapshot(template)
    if not template.is_active or project.summary_template_snapshot != expected_snapshot:
        raise TemplateStructureError("项目冻结汇总模板信息不一致")
    payload = read_template_bytes(template)
    _workbook, manifest, signature = _load_validated_workbook(payload)
    if signature != template.structure_signature or manifest != template.structure_manifest:
        raise TemplateStructureError("冻结汇总模板结构签名不一致")
    return template, payload


def _structure_digest(manifest):
    return _sha256(
        json.dumps(
            manifest,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
            default=str,
        ).encode("utf-8")
    )


def _validated_export_rows(project, template):
    project_subjects = list(
        project.subjects.select_related("subject").order_by("public_id")
    )
    results = list(
        project.aggregate_results.select_related("project_subject").order_by(
            "project_subject__public_id", "pk"
        )
    )
    results_by_subject = {}
    for result in results:
        project_subject = result.project_subject
        if (
            result.project_id != project.pk
            or project_subject.project_id != project.pk
            or result.subject_id != project_subject.subject_id
            or str(result.frozen_subject_public_id)
            != project_subject.subject_snapshot.get("public_id")
            or result.frozen_subject_name
            != project_subject.subject_snapshot.get("name")
            or project_subject.pk in results_by_subject
        ):
            raise TemplateStructureError(
                "汇总结果与项目冻结数据不一致", "REPORT_DATA_INVALID"
            )
        results_by_subject[project_subject.pk] = result
    rows = []
    employee_nos = set()
    for project_subject in project_subjects:
        snapshot = project_subject.reporting_snapshot
        if (
            not isinstance(snapshot, dict)
            or set(snapshot)
            != {"employee_no", "corporate_email", "department", "manager_name"}
            or not isinstance(project_subject.subject_snapshot, dict)
            or str(project_subject.subject.public_id)
            != project_subject.subject_snapshot.get("public_id")
        ):
            raise TemplateStructureError(
                "项目导出冻结数据无效", "REPORT_DATA_INVALID"
            )
        employee_no = snapshot.get("employee_no")
        if not isinstance(employee_no, str) or not employee_no or employee_no in employee_nos:
            raise TemplateStructureError(
                "项目导出冻结员工编号无效", "REPORT_DATA_INVALID"
            )
        employee_nos.add(employee_no)
        rows.append((project_subject, results_by_subject.get(project_subject.pk)))
    template_rows = set(template.structure_manifest.get("target_rows", {}))
    if employee_nos != template_rows:
        raise TemplateStructureError(
            "汇总模板目标行与项目冻结被评价人不一致", "REPORT_TARGET_ROWS_INVALID"
        )
    return rows


def _score_value(value):
    return None if value is None else float(value)


def export_summary_workbook(project):
    project = EvaluationProject.objects.select_related("summary_template").get(pk=project.pk)
    if project.status == EvaluationProject.Status.DRAFT or project.prepared_at is None:
        raise TemplateStructureError("项目尚未准备，不能导出汇总表")
    try:
        frozen_project_name(project)
    except ReportingSnapshotUnavailable as exc:
        raise TemplateStructureError(
            str(exc), "REPORTING_SNAPSHOT_UNAVAILABLE"
        ) from exc
    template, payload = verify_project_template_binding(project)
    workbook, original_manifest, original_signature = _load_validated_workbook(payload)
    if (
        original_manifest != template.structure_manifest
        or original_signature != template.structure_signature
    ):
        raise TemplateStructureError("冻结汇总模板结构签名不一致")
    rows = _validated_export_rows(project, template)
    summary = workbook["Sheet1"]
    target_rows = template.structure_manifest["target_rows"]
    try:
        for project_subject, result in rows:
            snapshot = project_subject.reporting_snapshot
            row = target_rows[snapshot["employee_no"]]
            values = {
                "A": safe_cell_text(snapshot["employee_no"]),
                "B": safe_cell_text(project_subject.subject_snapshot["name"]),
                "C": safe_cell_text(snapshot["corporate_email"]),
                "D": safe_cell_text(snapshot["department"]),
                "E": safe_cell_text(snapshot["manager_name"]),
                "I": _score_value(result.manager_score) if result else None,
                "J": _score_value(result.same_department_score) if result else None,
                "K": _score_value(result.cross_department_score) if result else None,
                "L": (
                    _score_value(result.total_score)
                    if result and result.status == result.Status.COMPLETE
                    else None
                ),
                "M": (
                    "完整"
                    if result and result.status == result.Status.COMPLETE
                    else "数据不完整"
                ),
            }
            for column, value in values.items():
                summary[f"{column}{row}"].value = value
    except OutputBoundaryError as exc:
        raise TemplateStructureError(str(exc), "REPORT_OUTPUT_BOUNDARY") from exc

    after_manifest = workbook_structure_manifest(workbook)
    if _structure_digest(after_manifest) != original_signature:
        raise TemplateStructureError("汇总导出改变了批准范围外的模板结构")
    output = BytesIO()
    workbook.save(output)
    artifact = output.getvalue()
    saved_workbook, saved_manifest, saved_signature = _load_validated_workbook(artifact)
    del saved_workbook
    if saved_manifest != after_manifest or saved_signature != original_signature:
        raise TemplateStructureError("汇总导出响应准备后模板结构不一致")
    return artifact
