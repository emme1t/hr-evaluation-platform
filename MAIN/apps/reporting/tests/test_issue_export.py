from io import BytesIO

import pytest
from openpyxl import load_workbook

from apps.reporting.services.issues import export_issue_workbook


pytestmark = pytest.mark.django_db


def _rows(payload):
    workbook = load_workbook(BytesIO(payload), data_only=False)
    worksheet = workbook["异常说明"]
    return workbook, list(worksheet.iter_rows(values_only=True))


def test_issue_workbook_has_stable_frozen_columns_and_remediation(incomplete_result):
    workbook, rows = _rows(export_issue_workbook(incomplete_result.project))

    assert workbook.sheetnames == ["异常说明"]
    assert rows[0] == (
        "错误码",
        "项目",
        "被评价人员工编号",
        "关系类型",
        "任务状态",
        "建议处理方式",
    )
    assert (
        "TASK_NOT_SUBMITTED",
        incomplete_result.project.name,
        incomplete_result.subject.employee_no,
        "manager",
        "pending",
        "联系评价人完成提交；截止后确认是否保持数据不完整",
    ) in rows
    assert any(row[0] == "INCOMPLETE_RESULT" for row in rows[1:])
    incomplete_row = next(row for row in rows[1:] if row[0] == "INCOMPLETE_RESULT")
    assert incomplete_row[2] == incomplete_result.subject.employee_no
    assert incomplete_row[3] == "manager"
    assert incomplete_row[4] == "数据不完整"


def test_issue_workbook_uses_frozen_employee_number_and_neutralizes_text(
    incomplete_result,
):
    project_subject = incomplete_result.project.subjects.get()
    snapshot = dict(project_subject.reporting_snapshot)
    snapshot["employee_no"] = "=FROZEN"
    project_subject.reporting_snapshot = snapshot
    project_subject.save(update_fields=["reporting_snapshot"])
    incomplete_result.subject.employee_no = "CURRENT-SECRET-PASSWORD"
    incomplete_result.subject.save(update_fields=["employee_no"])

    _workbook, rows = _rows(export_issue_workbook(incomplete_result.project))
    assert all(row[2] == "'=FROZEN" for row in rows[1:])
    searchable = "\n".join(str(value) for row in rows for value in row)
    assert "CURRENT-SECRET-PASSWORD" not in searchable


def test_issue_workbook_has_no_sensitive_fields_or_values(incomplete_result):
    incomplete_result.subject.wecom_userid = "cookie-authorization-token-value"
    incomplete_result.subject.save(update_fields=["wecom_userid"])
    _workbook, rows = _rows(export_issue_workbook(incomplete_result.project))
    lowered = "\n".join(str(value) for row in rows for value in row).lower()
    for forbidden in (
        "token",
        "cookie",
        "authorization",
        "password",
        "secret",
        "magic link",
        "cookie-authorization-token-value",
    ):
        assert forbidden not in lowered


def test_issue_export_rapid_requests_do_not_mutate_business_state(incomplete_result):
    project = incomplete_result.project
    before = (
        list(project.subjects.values()),
        list(project.tasks.values()),
        list(project.aggregate_results.values()),
    )
    for _ in range(3):
        assert export_issue_workbook(project)
    after = (
        list(project.subjects.values()),
        list(project.tasks.values()),
        list(project.aggregate_results.values()),
    )
    assert after == before
