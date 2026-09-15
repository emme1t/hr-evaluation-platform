import json
from io import BytesIO
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from apps.reporting.services.raw import RawExportError, export_raw_zip
from tests.factories import create_hr_user, create_user_with_role


pytestmark = pytest.mark.django_db


def _archive(payload):
    return ZipFile(BytesIO(payload))


def _workbook_values(payload):
    workbook = load_workbook(BytesIO(payload), data_only=False)
    return [
        cell.value
        for worksheet in workbook.worksheets
        for row in worksheet.iter_rows()
        for cell in row
        if cell.value is not None
    ]


def test_raw_export_is_deterministic_partitioned_and_manifested(
    project_results, hr_admin
):
    subject_id = project_results.subject.public_id
    first = export_raw_zip(project_results.project, subject_id, actor=hr_admin)
    second = export_raw_zip(project_results.project, subject_id, actor=hr_admin)

    assert first == second
    archive = _archive(first)
    names = archive.namelist()
    assert names[0] == "manifest.json"
    assert len(names) == 4
    assert all(
        name == "manifest.json" or name.startswith(f"{subject_id}/")
        for name in names
    )
    assert all(name.endswith(".xlsx") for name in names[1:])
    assert all(".." not in name and "\\" not in name for name in names)

    manifest = json.loads(archive.read("manifest.json"))
    assert manifest["project_public_id"] == str(project_results.project.public_id)
    assert manifest["subject_public_ids"] == [str(subject_id)]
    assert [entry["path"] for entry in manifest["members"]] == names[1:]
    assert all(len(entry["sha256"]) == 64 for entry in manifest["members"])


def test_raw_workbooks_use_only_frozen_submission_and_project_snapshots(
    project_results, hr_admin
):
    project_subject = project_results.project.subjects.get()
    frozen_subject_name = project_subject.subject_snapshot["name"]
    frozen_evaluator_names = {
        item["evaluator"]["name"]
        for item in project_subject.relationship_snapshot
    }
    project_results.subject.name = "Current Secret Password Subject"
    project_results.subject.employee_no = "CURRENT-SECRET"
    project_results.subject.save(update_fields=["name", "employee_no"])
    project_results.project.tasks.update(status="not_submitted")

    archive = _archive(
        export_raw_zip(project_results.project, actor=hr_admin)
    )
    all_values = []
    for name in archive.namelist()[1:]:
        values = _workbook_values(archive.read(name))
        all_values.extend(values)
        assert frozen_subject_name in values
        assert any(name in values for name in frozen_evaluator_names)
        assert "Current Secret Password Subject" not in values
        assert "CURRENT-SECRET" not in values
    assert "not_submitted" not in all_values


def test_raw_export_uses_frozen_project_name_and_remains_byte_deterministic(
    project_results, hr_admin
):
    frozen_name = project_results.project.reporting_snapshot["name"]
    before = export_raw_zip(project_results.project, actor=hr_admin)
    project_results.project.name = "Mutable Current Project Name"
    project_results.project.save(update_fields=["name"])
    after = export_raw_zip(project_results.project, actor=hr_admin)

    assert after == before
    archive = _archive(after)
    for name in archive.namelist()[1:]:
        values = _workbook_values(archive.read(name))
        assert frozen_name in values
        assert "Mutable Current Project Name" not in values


@pytest.mark.parametrize("role", ["HR_OPERATOR", "EVALUATOR"])
def test_raw_export_requires_hr_admin(role, project_results):
    actor = (
        create_hr_user(role)
        if role == "HR_OPERATOR"
        else create_user_with_role(role)
    )
    with pytest.raises(RawExportError) as raised:
        export_raw_zip(project_results.project, actor=actor)
    assert raised.value.code == "HR_ADMIN_REQUIRED"


def test_raw_subject_filter_rejects_cross_project_identifier_without_leak(
    project_results, incomplete_result, hr_admin
):
    foreign_id = incomplete_result.subject.public_id
    with pytest.raises(RawExportError) as raised:
        export_raw_zip(project_results.project, foreign_id, actor=hr_admin)
    assert raised.value.code == "SUBJECT_NOT_IN_PROJECT"
    assert str(foreign_id) not in str(raised.value)


@pytest.mark.parametrize("payload", ["=1+1", "+cmd", "-2+3", "@SUM(A1:A2)"])
def test_raw_export_neutralizes_cells_and_bounds_path_components(
    payload, project_results, hr_admin
):
    project_subject = project_results.project.subjects.get()
    snapshot = list(project_subject.relationship_snapshot)
    snapshot[0] = {
        **snapshot[0],
        "evaluator": {**snapshot[0]["evaluator"], "name": payload + "/" + "X" * 200},
    }
    project_subject.relationship_snapshot = snapshot
    project_subject.save(update_fields=["relationship_snapshot"])

    archive = _archive(export_raw_zip(project_results.project, actor=hr_admin))
    member = next(name for name in archive.namelist() if name.endswith(".xlsx"))
    assert len(member.encode("utf-8")) <= 240
    assert "\\" not in member and ".." not in member
    values = _workbook_values(archive.read(member))
    evaluator_value = next(value for value in values if isinstance(value, str) and payload in value)
    assert evaluator_value.startswith("'")


def test_raw_export_enforces_first_member_and_total_resource_boundaries(
    project_results, hr_admin, monkeypatch
):
    import apps.reporting.services.raw as raw

    monkeypatch.setattr(raw, "MAX_OUTPUT_MEMBERS", 2)
    with pytest.raises(RawExportError) as raised:
        export_raw_zip(project_results.project, actor=hr_admin)
    assert raised.value.code == "RAW_OUTPUT_MEMBER_LIMIT"

    monkeypatch.setattr(raw, "MAX_OUTPUT_MEMBERS", 100)
    monkeypatch.setattr(raw, "MAX_OUTPUT_TOTAL_BYTES", 100)
    with pytest.raises(RawExportError) as raised:
        export_raw_zip(project_results.project, actor=hr_admin)
    assert raised.value.code == "RAW_OUTPUT_BYTE_LIMIT"


def test_raw_export_contains_no_sensitive_key_or_current_secret_value(
    project_results, hr_admin
):
    project_results.subject.wecom_userid = "notification-magic-link-secret-value"
    project_results.subject.save(update_fields=["wecom_userid"])
    payload = export_raw_zip(project_results.project, actor=hr_admin)
    archive = _archive(payload)
    searchable = archive.read("manifest.json").decode("utf-8")
    for name in archive.namelist()[1:]:
        searchable += "\n" + "\n".join(map(str, _workbook_values(archive.read(name))))
    lowered = searchable.lower()
    for forbidden in (
        "token",
        "cookie",
        "authorization",
        "password",
        "secret",
        "magic link",
        "notification-magic-link-secret-value",
    ):
        assert forbidden not in lowered


def test_raw_export_redacts_sensitive_markers_already_present_in_frozen_text(
    project_results, hr_admin
):
    project_subject = project_results.project.subjects.get()
    template_snapshot = dict(project_subject.template_snapshot)
    items = [dict(item) for item in template_snapshot["items"]]
    items[0]["excellent_description"] = "password=fictional-credential"
    template_snapshot["items"] = items
    relationships = [dict(item) for item in project_subject.relationship_snapshot]
    relationships[0] = {
        **relationships[0],
        "evaluator": {
            **relationships[0]["evaluator"],
            "name": "Secret=fictional-credential",
        },
    }
    project_subject.template_snapshot = template_snapshot
    project_subject.relationship_snapshot = relationships
    project_subject.save(
        update_fields=["template_snapshot", "relationship_snapshot"]
    )

    archive = _archive(export_raw_zip(project_results.project, actor=hr_admin))
    searchable = archive.read("manifest.json").decode("utf-8")
    for name in archive.namelist()[1:]:
        searchable += "\n" + name
        searchable += "\n" + "\n".join(map(str, _workbook_values(archive.read(name))))
    lowered = searchable.lower()
    assert "password" not in lowered
    assert "secret" not in lowered
    assert "fictional-credential" not in lowered
    assert "已移除敏感内容" in searchable
