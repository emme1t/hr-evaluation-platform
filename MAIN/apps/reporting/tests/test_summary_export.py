from io import BytesIO

import pytest
from openpyxl import load_workbook

from apps.evaluations.models import AggregateResult
from apps.reporting.services.summary import (
    TemplateStructureError,
    export_summary_workbook,
)
from apps.reporting.tests.workbook_helpers import find_employee_row
from apps.reporting.services.common import OutputBoundaryError, safe_cell_text
from apps.roster.models import Employee


pytestmark = pytest.mark.django_db


def _independent_structure(workbook):
    return {
        "sheetnames": workbook.sheetnames,
        "sheets": {
            sheet.title: {
                "dimension": sheet.calculate_dimension(),
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "merged": tuple(sorted(map(str, sheet.merged_cells.ranges))),
                "columns": {
                    key: (value.width, value.hidden, value.outlineLevel)
                    for key, value in sheet.column_dimensions.items()
                },
                "styles": {
                    cell.coordinate: cell._style
                    for cell in sheet._cells.values()
                },
                "formulas": {
                    cell.coordinate: cell.value
                    for cell in sheet._cells.values()
                    if cell.data_type == "f"
                },
                "tables": {
                    table.name: (
                        table.ref,
                        tuple(
                            (
                                column.name,
                                getattr(column.calculatedColumnFormula, "text", None),
                                getattr(column.totalsRowFormula, "text", None),
                            )
                            for column in table.tableColumns
                        ),
                        table.tableStyleInfo.name if table.tableStyleInfo else None,
                    )
                    for table in sheet.tables.values()
                },
            }
            for sheet in workbook.worksheets
        },
        "defined_names": tuple(
            sorted(
                (item.name, item.attr_text, item.localSheetId)
                for item in workbook.defined_names.values()
            )
        ),
    }


def test_summary_export_changes_only_approved_values_and_preserves_exact_structure(
    project_results,
):
    before = load_workbook(BytesIO(project_results.template_bytes), data_only=False)
    artifact = export_summary_workbook(project_results.project)
    after = load_workbook(BytesIO(artifact), data_only=False)

    assert _independent_structure(after) == _independent_structure(before)
    row = find_employee_row(after["Sheet1"], project_results.subject.employee_no)
    assert after["Sheet1"][f"B{row}"].value == project_results.frozen_name
    assert after["Sheet1"][f"I{row}"].value == 5
    assert after["Sheet1"][f"J{row}"].value == 3.6
    assert after["Sheet1"][f"K{row}"].value == 3
    assert after["Sheet1"][f"L{row}"].value == 4.18
    assert after["Sheet1"][f"M{row}"].value == "完整"
    assert after["Sheet3"]["A1"].value == "=SUM(1,2)"

    target_cells = {
        f"{column}{row}"
        for column in ("A", "B", "C", "D", "E", "I", "J", "K", "L", "M")
    }
    for worksheet in before.worksheets:
        for coordinate, cell in worksheet._cells.items():
            after_cell = after[worksheet.title]._cells[coordinate]
            if worksheet.title != "Sheet1" or cell.coordinate not in target_cells:
                assert after_cell.value == cell.value


def test_incomplete_employee_has_blank_total_without_group_renormalization(
    incomplete_result,
):
    workbook = load_workbook(
        BytesIO(export_summary_workbook(incomplete_result.project)), data_only=False
    )
    row = find_employee_row(workbook["Sheet1"], incomplete_result.subject.employee_no)

    assert workbook["Sheet1"][f"I{row}"].value is None
    assert workbook["Sheet1"][f"J{row}"].value == 3.6
    assert workbook["Sheet1"][f"K{row}"].value == 3
    assert workbook["Sheet1"][f"L{row}"].value is None
    assert workbook["Sheet1"][f"M{row}"].value == "数据不完整"


@pytest.mark.parametrize("payload", ["=1+1", "+cmd", "-2+3", "@SUM(A1:A2)"])
def test_summary_export_neutralizes_formula_injection_payloads(
    payload, project_results
):
    project_subject = project_results.project.subjects.get()
    reporting_snapshot = dict(project_subject.reporting_snapshot)
    reporting_snapshot["manager_name"] = payload
    project_subject.reporting_snapshot = reporting_snapshot
    project_subject.save(update_fields=["reporting_snapshot"])

    workbook = load_workbook(
        BytesIO(export_summary_workbook(project_results.project)), data_only=False
    )
    row = find_employee_row(workbook["Sheet1"], project_results.subject.employee_no)
    assert workbook["Sheet1"][f"E{row}"].value == f"'{payload}"
    assert workbook["Sheet1"][f"E{row}"].data_type == "s"


def test_summary_export_accepts_32767_characters_and_rejects_32768(
    project_results,
):
    project_subject = project_results.project.subjects.get()
    snapshot = dict(project_subject.reporting_snapshot)
    snapshot["manager_name"] = "X" * 32767
    project_subject.reporting_snapshot = snapshot
    project_subject.save(update_fields=["reporting_snapshot"])
    workbook = load_workbook(BytesIO(export_summary_workbook(project_results.project)))
    row = find_employee_row(workbook["Sheet1"], project_results.subject.employee_no)
    assert len(workbook["Sheet1"][f"E{row}"].value) == 32767

    snapshot["manager_name"] = "X" * 32768
    project_subject.reporting_snapshot = snapshot
    project_subject.save(update_fields=["reporting_snapshot"])
    with pytest.raises(TemplateStructureError, match="32767"):
        export_summary_workbook(project_results.project)


@pytest.mark.parametrize("prefix", ["=", "+", "-", "@"])
@pytest.mark.parametrize(
    ("input_length", "accepted"),
    [(32766, True), (32767, False), (32768, False)],
)
def test_formula_neutralization_enforces_final_cell_length_without_truncation(
    prefix, input_length, accepted
):
    payload = prefix + "X" * (input_length - 1)
    if accepted:
        result = safe_cell_text(payload)
        assert result == "'" + payload
        assert len(result) == 32767
    else:
        with pytest.raises(OutputBoundaryError, match="32767"):
            safe_cell_text(payload)


def test_sensitive_redaction_is_applied_before_final_cell_length_check():
    assert safe_cell_text("password=" + "X" * 32768) == "已移除敏感内容"


def test_summary_export_rejects_cross_project_result_subject_without_leaking(
    project_results,
):
    other = Employee.objects.exclude(pk=project_results.subject.pk).first()
    assert other is not None
    AggregateResult.objects.filter(pk=project_results.result.pk).update(subject=other)

    with pytest.raises(TemplateStructureError) as raised:
        export_summary_workbook(project_results.project)
    assert raised.value.code == "REPORT_DATA_INVALID"
    assert other.employee_no not in str(raised.value)


def test_summary_export_rejects_frozen_template_hash_or_structure_tampering(
    project_results,
):
    project = project_results.project
    snapshot = dict(project.summary_template_snapshot)
    snapshot["sha256"] = "0" * 64
    project.summary_template_snapshot = snapshot
    project.save(update_fields=["summary_template_snapshot"])
    with pytest.raises(TemplateStructureError, match="冻结"):
        export_summary_workbook(project)


def test_rapid_summary_exports_do_not_mutate_business_state(project_results):
    project = project_results.project
    before = {
        "project": (project.status, project.rule_snapshot, project.prepared_at),
        "subjects": list(project.subjects.values()),
        "tasks": list(project.tasks.values()),
        "results": list(project.aggregate_results.values()),
    }

    artifacts = [export_summary_workbook(project) for _ in range(3)]

    after = {
        "project": tuple(
            type(project).objects.filter(pk=project.pk).values_list(
                "status", "rule_snapshot", "prepared_at"
            ).get()
        ),
        "subjects": list(project.subjects.values()),
        "tasks": list(project.tasks.values()),
        "results": list(project.aggregate_results.values()),
    }
    assert all(artifacts)
    assert after == before
