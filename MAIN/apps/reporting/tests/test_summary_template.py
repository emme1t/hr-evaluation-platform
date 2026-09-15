from io import BytesIO
from copy import copy
from concurrent.futures import ThreadPoolExecutor
import struct
import warnings
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

import pytest
from django.core.exceptions import ValidationError
from django.contrib.auth import get_user_model
from django.db import close_old_connections
from openpyxl import load_workbook

from apps.reporting.services.summary import (
    SummaryTemplateError,
    _load_validated_workbook,
    assign_summary_template,
    register_summary_template,
)
from apps.reporting.storage import PrivateStorageError, private_root
from apps.reporting.models import SummaryWorkbookTemplate
from apps.reporting.tests.workbook_helpers import build_summary_template_bytes
from tests.factories import create_project


@pytest.mark.django_db
def test_summary_template_is_versioned_private_immutable_and_admin_assigned(
    draft_project, hr_admin, settings, tmp_path
):
    settings.REPORTING_PRIVATE_ROOT = tmp_path / "reporting-private"
    payload = build_summary_template_bytes(("P001",))

    first = register_summary_template(payload, "summary-template.xlsx", hr_admin)
    second = register_summary_template(payload, "summary-template.xlsx", hr_admin)
    assign_summary_template(draft_project, second, hr_admin)

    draft_project.refresh_from_db()
    assert (first.version, second.version) == (1, 2)
    assert len(second.sha256) == 64
    assert len(second.structure_signature) == 64
    assert not second.storage_path.startswith(("/", "static/", "media/"))
    assert (settings.REPORTING_PRIVATE_ROOT / second.storage_path).read_bytes() == payload
    assert draft_project.summary_template_id == second.id
    assert draft_project.summary_template_snapshot == {
        "public_id": str(second.public_id),
        "version": 2,
        "sha256": second.sha256,
        "structure_signature": second.structure_signature,
    }

    second.sha256 = "0" * 64
    with pytest.raises(ValidationError):
        second.save()
    with pytest.raises(ValidationError):
        second.delete()
    with pytest.raises(ValidationError):
        type(second).objects.filter(pk=second.pk).update(is_active=False)
    with pytest.raises(ValidationError):
        type(second).objects.filter(pk=second.pk).delete()


@pytest.mark.django_db
def test_summary_template_assignment_is_admin_only_and_draft_only(
    draft_project, hr_admin, settings, tmp_path
):
    settings.REPORTING_PRIVATE_ROOT = tmp_path / "reporting-private"
    template = register_summary_template(
        build_summary_template_bytes(("P001",)), "summary.xlsx", hr_admin
    )
    operator = __import__("tests.factories", fromlist=["create_hr_user"]).create_hr_user(
        "HR_OPERATOR"
    )

    with pytest.raises(SummaryTemplateError, match="管理员"):
        assign_summary_template(draft_project, template, operator)

    draft_project.status = "ready"
    draft_project.save(update_fields=["status"])
    with pytest.raises(SummaryTemplateError, match="草稿"):
        assign_summary_template(draft_project, template, hr_admin)


@pytest.mark.django_db
def test_prepare_requires_verified_binding_and_freezes_reporting_identity(
    draft_project, hr_admin
):
    from apps.evaluations.services.projects import prepare_project

    expected_template_snapshot = dict(draft_project.summary_template_snapshot)
    prepare_project(draft_project, hr_admin)
    draft_project.refresh_from_db()
    project_subject = draft_project.subjects.get()

    assert draft_project.summary_template_snapshot == expected_template_snapshot
    assert draft_project.reporting_snapshot == {"name": draft_project.name}
    assert project_subject.reporting_snapshot == {
        "employee_no": "P001",
        "corporate_email": "p001@example.test",
        "department": "Test Center/Test Team",
        "manager_name": "Project Manager",
    }


@pytest.mark.django_db
def test_prepare_without_summary_template_fails_closed_before_business_mutation(
    hr_admin,
):
    from apps.evaluations.services.projects import prepare_project

    project = create_project(bind_summary_template=False)
    with pytest.raises(SummaryTemplateError, match="绑定"):
        prepare_project(project, hr_admin)
    project.refresh_from_db()
    assert project.status == "draft"
    assert project.tasks.count() == 0
    assert all(
        not item.reporting_snapshot and not item.subject_snapshot
        for item in project.subjects.all()
    )


@pytest.mark.django_db
@pytest.mark.parametrize("filename", ["summary.xls", "summary.xlsm", "summary.csv"])
def test_registration_rejects_unsupported_extensions(
    filename, hr_admin, settings, tmp_path
):
    settings.REPORTING_PRIVATE_ROOT = tmp_path / "reporting-private"
    with pytest.raises(SummaryTemplateError, match="xlsx"):
        register_summary_template(build_summary_template_bytes(), filename, hr_admin)


def _rewrite_zip(payload, mutate):
    source = ZipFile(BytesIO(payload))
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            replacement = mutate(info, data)
            if replacement is None:
                target.writestr(info, data)
            else:
                replacement_info, replacement_data = replacement
                target.writestr(replacement_info, replacement_data)
    return output.getvalue()


def _mark_first_member_encrypted(payload):
    changed = bytearray(payload)
    local = changed.find(b"PK\x03\x04")
    central = changed.find(b"PK\x01\x02")
    assert local >= 0 and central >= 0
    changed[local + 6 : local + 8] = (
        int.from_bytes(changed[local + 6 : local + 8], "little") | 1
    ).to_bytes(2, "little")
    changed[central + 8 : central + 10] = (
        int.from_bytes(changed[central + 8 : central + 10], "little") | 1
    ).to_bytes(2, "little")
    return bytes(changed)


def _archive_member_offsets(content, name):
    encoded_name = name.encode("utf-8")
    offsets = {}
    for key, signature, fixed_size, name_length_offset in (
        ("local", b"PK\x03\x04", 30, 26),
        ("central", b"PK\x01\x02", 46, 28),
    ):
        cursor = 0
        while True:
            offset = content.find(signature, cursor)
            if offset < 0:
                raise AssertionError(f"missing {key} header for {name}")
            name_length = struct.unpack_from(
                "<H", content, offset + name_length_offset
            )[0]
            candidate = content[offset + fixed_size : offset + fixed_size + name_length]
            if candidate == encoded_name:
                offsets[key] = offset
                break
            cursor = offset + 4
    return offsets


def _corrupt_member_crc(content, name):
    changed = bytearray(content)
    offsets = _archive_member_offsets(changed, name)
    struct.pack_into("<I", changed, offsets["local"] + 14, 0)
    struct.pack_into("<I", changed, offsets["central"] + 16, 0)
    return bytes(changed)


@pytest.mark.django_db
@pytest.mark.parametrize(
    ("label", "mutator", "message"),
    [
        (
            "traversal",
            lambda info, data: (
                (ZipInfo("../escape.xml"), b"<x/>")
                if info.filename == "[Content_Types].xml"
                else None
            ),
            "路径",
        ),
        (
            "malformed-relationship",
            lambda info, data: (
                (info, b"<Relationships><Relationship")
                if info.filename == "_rels/.rels"
                else None
            ),
            "关系",
        ),
        (
            "external-link",
            lambda info, data: (
                (ZipInfo("xl/externalLinks/externalLink1.xml"), b"<x/>")
                if info.filename == "[Content_Types].xml"
                else None
            ),
            "外部链接",
        ),
        (
            "macro",
            lambda info, data: (
                (ZipInfo("xl/vbaProject.bin"), b"macro")
                if info.filename == "[Content_Types].xml"
                else None
            ),
            "宏",
        ),
    ],
)
def test_registration_rejects_unsafe_container_members(
    label, mutator, message, hr_admin, settings, tmp_path
):
    settings.REPORTING_PRIVATE_ROOT = tmp_path / "reporting-private"
    payload = _rewrite_zip(build_summary_template_bytes(), mutator)
    with pytest.raises(SummaryTemplateError, match=message):
        register_summary_template(payload, "summary.xlsx", hr_admin)


@pytest.mark.django_db
def test_registration_rejects_encrypted_member(hr_admin, settings, tmp_path):
    settings.REPORTING_PRIVATE_ROOT = tmp_path / "reporting-private"
    payload = _mark_first_member_encrypted(build_summary_template_bytes())
    with pytest.raises(SummaryTemplateError, match="加密"):
        register_summary_template(payload, "summary.xlsx", hr_admin)


@pytest.mark.django_db
def test_registration_rejects_duplicate_member_and_first_resource_boundary(
    hr_admin, settings, tmp_path
):
    settings.REPORTING_PRIVATE_ROOT = tmp_path / "reporting-private"
    payload = build_summary_template_bytes()
    source = ZipFile(BytesIO(payload))
    duplicate = BytesIO()
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        with ZipFile(duplicate, "w", ZIP_DEFLATED) as archive:
            for info in source.infolist():
                archive.writestr(info, source.read(info.filename))
            archive.writestr("xl/workbook.xml", source.read("xl/workbook.xml"))
    with pytest.raises(SummaryTemplateError, match="重复"):
        register_summary_template(duplicate.getvalue(), "summary.xlsx", hr_admin)

    with pytest.raises(SummaryTemplateError, match="10MB"):
        register_summary_template(b"x" * (10 * 1024 * 1024 + 1), "summary.xlsx", hr_admin)


@pytest.mark.django_db
@pytest.mark.parametrize(
    ("mutate", "message"),
    [
        (lambda wb: setattr(wb["Sheet3"]["A1"], "value", "=WEBSERVICE(\"https://example.test\")"), "不安全公式"),
        (lambda wb: wb.defined_names.add(__import__("openpyxl").workbook.defined_name.DefinedName("Auto_Open", attr_text="=EXEC(\"x\")")), "XLM"),
        (lambda wb: setattr(wb["Sheet1"]["A1048576"], "value", "SPARSE"), "资源"),
    ],
)
def test_registration_rejects_unsafe_workbook_content(
    mutate, message, hr_admin, settings, tmp_path
):
    settings.REPORTING_PRIVATE_ROOT = tmp_path / "reporting-private"
    workbook = load_workbook(BytesIO(build_summary_template_bytes()))
    mutate(workbook)
    output = BytesIO()
    workbook.save(output)
    with pytest.raises(SummaryTemplateError, match=message):
        register_summary_template(output.getvalue(), "summary.xlsx", hr_admin)


@pytest.mark.django_db
def test_registration_rejects_32768_character_xml_cell(
    hr_admin, settings, tmp_path
):
    settings.REPORTING_PRIVATE_ROOT = tmp_path / "reporting-private"

    def mutate(info, data):
        if info.filename != "xl/worksheets/sheet1.xml":
            return None
        assert b"P001" in data
        return info, data.replace(b"P001", b"X" * 32768, 1)

    payload = _rewrite_zip(build_summary_template_bytes(), mutate)
    with pytest.raises(SummaryTemplateError, match="32767"):
        register_summary_template(payload, "summary.xlsx", hr_admin)


@pytest.mark.django_db
def test_registration_rejects_macro_enabled_content_type_without_vba_member(
    hr_admin, settings, tmp_path
):
    settings.REPORTING_PRIVATE_ROOT = tmp_path / "reporting-private"

    def mutate(info, data):
        if info.filename != "[Content_Types].xml":
            return None
        return info, data.replace(
            b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
            b"application/vnd.ms-excel.sheet.macroEnabled.main+xml",
        )

    payload = _rewrite_zip(build_summary_template_bytes(), mutate)
    with pytest.raises(SummaryTemplateError, match="宏"):
        register_summary_template(payload, "summary.xlsx", hr_admin)


@pytest.mark.django_db
def test_registration_rejects_compression_ratio_bomb_before_workbook_parse(
    hr_admin, settings, tmp_path
):
    settings.REPORTING_PRIVATE_ROOT = tmp_path / "reporting-private"
    source = ZipFile(BytesIO(build_summary_template_bytes()))
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED, compresslevel=9) as archive:
        archive.writestr("bomb.bin", b"0" * (2 * 1024 * 1024))
        for info in source.infolist():
            archive.writestr(info, source.read(info.filename))
    with pytest.raises(SummaryTemplateError, match="压缩比"):
        register_summary_template(output.getvalue(), "summary.xlsx", hr_admin)


@pytest.mark.django_db
def test_registered_template_bulk_apis_preserve_immutability_but_allow_pure_inserts(
    hr_admin, settings, tmp_path
):
    settings.REPORTING_PRIVATE_ROOT = tmp_path / "reporting-private"
    registered = register_summary_template(
        build_summary_template_bytes(), "immutable.xlsx", hr_admin
    )

    registered.sha256 = "0" * 64
    with pytest.raises(ValidationError, match="不能修改"):
        SummaryWorkbookTemplate.objects.bulk_update([registered], ["sha256"])

    conflict = SummaryWorkbookTemplate(
        filename=registered.filename,
        version=registered.version,
        storage_path="templates/conflict.xlsx",
        sha256="1" * 64,
        structure_signature=registered.structure_signature,
        structure_manifest=registered.structure_manifest,
        created_by=hr_admin,
    )
    with pytest.raises(ValidationError, match="不能修改"):
        SummaryWorkbookTemplate.objects.bulk_create(
            [conflict],
            update_conflicts=True,
            update_fields=["sha256"],
            unique_fields=["filename", "version"],
        )

    pure_insert = SummaryWorkbookTemplate(
        filename="pure-insert.xlsx",
        version=1,
        storage_path="templates/pure-insert.xlsx",
        sha256="2" * 64,
        structure_signature=registered.structure_signature,
        structure_manifest=registered.structure_manifest,
        created_by=hr_admin,
    )
    SummaryWorkbookTemplate.objects.bulk_create([pure_insert])
    assert SummaryWorkbookTemplate.objects.filter(pk=pure_insert.pk).exists()


@pytest.mark.django_db
def test_registration_reads_non_xml_member_and_rejects_crc_failure(
    hr_admin, settings, tmp_path
):
    settings.REPORTING_PRIVATE_ROOT = tmp_path / "reporting-private"
    output = BytesIO(build_summary_template_bytes())
    member_name = "xl/media/fictional.bin"
    with ZipFile(output, "a", ZIP_DEFLATED) as archive:
        archive.writestr(member_name, b"fictional-non-xml-payload")
    payload = _corrupt_member_crc(output.getvalue(), member_name)

    with pytest.raises(SummaryTemplateError, match="安全读取"):
        register_summary_template(payload, "summary.xlsx", hr_admin)


@pytest.mark.django_db
def test_registration_rejects_unsafe_defined_name_without_equals_prefix(
    hr_admin, settings, tmp_path
):
    from openpyxl.workbook.defined_name import DefinedName

    settings.REPORTING_PRIVATE_ROOT = tmp_path / "reporting-private"
    workbook = load_workbook(BytesIO(build_summary_template_bytes()))
    workbook.defined_names.add(
        DefinedName("UnsafeName", attr_text='WEBSERVICE("https://example.test")')
    )
    output = BytesIO()
    workbook.save(output)

    with pytest.raises(SummaryTemplateError, match="不安全公式"):
        register_summary_template(output.getvalue(), "summary.xlsx", hr_admin)


@pytest.mark.django_db
def test_registration_rejects_remote_image_in_ordinary_cell_formula(
    hr_admin, settings, tmp_path
):
    settings.REPORTING_PRIVATE_ROOT = tmp_path / "reporting-private"
    workbook = load_workbook(BytesIO(build_summary_template_bytes()))
    workbook["Sheet3"]["A1"] = '=_xlfn.IMAGE("https://example.test/pixel.png")'
    output = BytesIO()
    workbook.save(output)

    with pytest.raises(SummaryTemplateError, match="不安全公式"):
        register_summary_template(output.getvalue(), "summary.xlsx", hr_admin)


@pytest.mark.django_db
def test_registration_rejects_remote_image_in_array_formula_object(
    hr_admin, settings, tmp_path
):
    settings.REPORTING_PRIVATE_ROOT = tmp_path / "reporting-private"

    def mutate(info, data):
        if info.filename != "xl/worksheets/sheet3.xml":
            return None
        assert b"<f>SUM(1,2)</f>" in data
        return info, data.replace(
            b"<f>SUM(1,2)</f>",
            b'<f t="array" ref="A1">IMAGE("https://example.test/pixel.png")</f>',
            1,
        )

    payload = _rewrite_zip(build_summary_template_bytes(), mutate)
    with pytest.raises(SummaryTemplateError, match="不安全公式"):
        register_summary_template(payload, "summary.xlsx", hr_admin)


@pytest.mark.django_db
def test_registration_rejects_remote_image_in_defined_name_without_equals_prefix(
    hr_admin, settings, tmp_path
):
    from openpyxl.workbook.defined_name import DefinedName

    settings.REPORTING_PRIVATE_ROOT = tmp_path / "reporting-private"
    workbook = load_workbook(BytesIO(build_summary_template_bytes()))
    workbook.defined_names.add(
        DefinedName(
            "RemoteImage",
            attr_text='IMAGE("https://example.test/pixel.png")',
        )
    )
    output = BytesIO()
    workbook.save(output)

    with pytest.raises(SummaryTemplateError, match="不安全公式"):
        register_summary_template(output.getvalue(), "summary.xlsx", hr_admin)


@pytest.mark.django_db
def test_registration_rejects_unsafe_array_formula_object(
    hr_admin, settings, tmp_path
):
    settings.REPORTING_PRIVATE_ROOT = tmp_path / "reporting-private"

    def mutate(info, data):
        if info.filename != "xl/worksheets/sheet3.xml":
            return None
        assert b"<f>SUM(1,2)</f>" in data
        return info, data.replace(
            b"<f>SUM(1,2)</f>",
            b'<f t="array" ref="A1">WEBSERVICE("https://example.test")</f>',
            1,
        )

    payload = _rewrite_zip(build_summary_template_bytes(), mutate)
    with pytest.raises(SummaryTemplateError, match="不安全公式"):
        register_summary_template(payload, "summary.xlsx", hr_admin)


@pytest.mark.django_db
def test_registration_rejects_raw_worksheet_declared_sparse_dimension(
    hr_admin, settings, tmp_path
):
    settings.REPORTING_PRIVATE_ROOT = tmp_path / "reporting-private"

    def mutate(info, data):
        if info.filename != "xl/worksheets/sheet1.xml":
            return None
        assert b'ref="A1:M2"' in data
        return info, data.replace(b'ref="A1:M2"', b'ref="A1:XFD1048576"', 1)

    payload = _rewrite_zip(build_summary_template_bytes(), mutate)
    with pytest.raises(SummaryTemplateError, match="稀疏维度"):
        register_summary_template(payload, "summary.xlsx", hr_admin)


@pytest.mark.django_db
def test_structure_signature_uses_semantic_style_definitions(
    hr_admin, settings, tmp_path
):
    settings.REPORTING_PRIVATE_ROOT = tmp_path / "reporting-private"
    original = build_summary_template_bytes()
    workbook = load_workbook(BytesIO(original))
    changed_font = copy(workbook["Sheet2"]["A1"].font)
    changed_font.color = "CC0000"
    workbook["Sheet2"]["A1"].font = changed_font
    changed = BytesIO()
    workbook.save(changed)

    _original_workbook, original_manifest, original_signature = _load_validated_workbook(
        original
    )
    _changed_workbook, changed_manifest, changed_signature = _load_validated_workbook(
        changed.getvalue()
    )
    assert original_manifest != changed_manifest
    assert original_signature != changed_signature


def test_private_root_is_explicit_external_and_rejects_repository_paths(
    settings, tmp_path
):
    settings.REPORTING_PRIVATE_ROOT = ""
    with pytest.raises(PrivateStorageError, match="配置"):
        private_root()

    settings.REPORTING_PRIVATE_ROOT = settings.BASE_DIR / "private" / "reporting"
    with pytest.raises(PrivateStorageError, match="仓库"):
        private_root()

    settings.REPORTING_PRIVATE_ROOT = tmp_path / "external-reporting"
    assert private_root() == (tmp_path / "external-reporting").resolve()


def test_upload_reader_rejects_declared_and_streamed_overflow_without_read(
    monkeypatch,
):
    import apps.reporting.views as views

    monkeypatch.setattr(views, "MAX_TEMPLATE_BYTES", 4)

    class ChunkedUpload:
        name = "summary.xlsx"

        def __init__(self, size, chunks):
            self.size = size
            self._chunks = chunks
            self.chunk_calls = 0

        def read(self, *args, **kwargs):
            raise AssertionError("upload view must not use unbounded read()")

        def chunks(self, *args, **kwargs):
            for chunk in self._chunks:
                self.chunk_calls += 1
                yield chunk

    declared_oversize = ChunkedUpload(5, [b"must-not-be-read"])
    with pytest.raises(SummaryTemplateError, match="10MB"):
        views._read_uploaded_template(declared_oversize)
    assert declared_oversize.chunk_calls == 0

    streamed_oversize = ChunkedUpload(4, [b"abcd", b"e", b"must-not-be-read"])
    with pytest.raises(SummaryTemplateError, match="10MB"):
        views._read_uploaded_template(streamed_oversize)
    assert streamed_oversize.chunk_calls == 2

    accepted = ChunkedUpload(4, [b"ab", b"cd"])
    assert views._read_uploaded_template(accepted) == b"abcd"
    assert accepted.chunk_calls == 2


@pytest.mark.django_db(transaction=True)
def test_concurrent_first_registration_has_stable_conflict_and_no_orphan(
    hr_admin, settings, tmp_path
):
    settings.REPORTING_PRIVATE_ROOT = tmp_path / "reporting-private"
    payload = build_summary_template_bytes()
    def register_once():
        close_old_connections()
        try:
            actor = get_user_model().objects.get(pk=hr_admin.pk)
            try:
                registered = register_summary_template(
                    payload, "concurrent.xlsx", actor
                )
                return ("ok", registered.version)
            except Exception as exc:
                return ("error", exc)
        finally:
            close_old_connections()

    with ThreadPoolExecutor(max_workers=2) as pool:
        outcomes = [future.result(timeout=20) for future in [
            pool.submit(register_once),
            pool.submit(register_once),
        ]]

    successes = [value for status, value in outcomes if status == "ok"]
    failures = [value for status, value in outcomes if status == "error"]
    assert sorted(successes) in ([1], [1, 2]), repr(outcomes)
    assert all(
        isinstance(error, SummaryTemplateError)
        and error.code == "SUMMARY_TEMPLATE_REGISTRATION_CONFLICT"
        for error in failures
    )
    database_paths = {
        row.storage_path
        for row in SummaryWorkbookTemplate.objects.filter(filename="concurrent.xlsx")
    }
    disk_paths = {
        path.relative_to(settings.REPORTING_PRIVATE_ROOT).as_posix()
        for path in settings.REPORTING_PRIVATE_ROOT.rglob("*.xlsx")
    }
    assert disk_paths == database_paths
