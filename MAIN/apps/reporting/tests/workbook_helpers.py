from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table, TableFormula, TableStyleInfo


SUMMARY_HEADERS = {
    "A": "员工编号",
    "B": "被评价人",
    "C": "邮箱",
    "D": "部门",
    "E": "上级",
    "I": "上级均分",
    "J": "同部门均分",
    "K": "跨部门均分",
    "L": "总分",
    "M": "状态",
}


def build_summary_template_bytes(employee_nos=("P001",)):
    """Build a fictional workbook in memory; never reads a desktop template."""
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Sheet1"
    for column, header in SUMMARY_HEADERS.items():
        summary[f"{column}1"] = header
        summary[f"{column}1"].font = Font(bold=True)
        summary[f"{column}1"].fill = PatternFill("solid", fgColor="D9EAF7")
    for index, employee_no in enumerate(employee_nos, start=2):
        summary[f"A{index}"] = employee_no
        for column_index in range(1, 14):
            summary.cell(index, column_index).number_format = "@"
    for column_index in range(1, 14):
        summary.column_dimensions[get_column_letter(column_index)].width = 14
    summary.column_dimensions["B"].width = 20
    summary.freeze_panes = "A2"

    notes = workbook.create_sheet("Sheet2")
    notes["A1"] = "虚构汇总说明"
    notes["A1"].font = Font(italic=True, color="336699")
    notes.merge_cells("A1:C1")
    notes.append(())
    notes.append(("项目", "虚构分值"))
    notes.append(("甲", 2))
    notes.append(("乙", 3))
    table = Table(displayName="FictionalSummaryTable", ref="A3:B5")
    table._initialise_columns()
    table.tableColumns[0].name = "项目"
    table.tableColumns[1].name = "虚构分值"
    table.tableColumns[1].calculatedColumnFormula = TableFormula(
        attr_text="SUM([虚构分值])"
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    notes.add_table(table)
    workbook.defined_names.add(
        DefinedName("FictionalRange", attr_text="'Sheet2'!$A$3:$B$5")
    )

    formulas = workbook.create_sheet("Sheet3")
    formulas["A1"] = "=SUM(1,2)"
    formulas["B1"] = "保留公式与样式"
    formulas.column_dimensions["A"].width = 18

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def find_employee_row(worksheet, employee_no):
    rows = [
        row
        for row in range(2, worksheet.max_row + 1)
        if worksheet[f"A{row}"].value == employee_no
    ]
    if len(rows) != 1:
        raise AssertionError(
            f"expected one row for fictional employee {employee_no!r}, got {rows}"
        )
    return rows[0]
