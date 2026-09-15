from __future__ import annotations

from dataclasses import dataclass
import csv
from hashlib import sha256
from io import BytesIO, StringIO
import json
from pathlib import Path
import struct
import zlib
from zipfile import BadZipFile, ZIP_DEFLATED, ZIP_STORED, ZipFile, ZipInfo

from django.db import IntegrityError, transaction
from django.db.models import Q
from django.utils import timezone

from apps.audit.services import record_audit
from openpyxl import load_workbook

from .models import Employee, EmployeeCategory, EvaluationRelationship, ImportBatch, ImportIssue
from .services import (
    RosterValidationError,
    create_employee_record,
    normalize_employee_fields,
    require_hr_actor,
    update_employee_record,
    upsert_relationship,
    validate_relationship_entities,
)


MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_XLSX_UNCOMPRESSED_BYTES = 20 * 1024 * 1024
MAX_XLSX_MEMBERS = 1000
MAX_IMPORT_ROWS = 500
MAX_IMPORT_ISSUES = 100
XLSX_STREAM_CHUNK_BYTES = 64 * 1024
ALLOWED_EXTENSIONS = {".xlsx", ".csv"}
ROSTER_HEADERS = (
    "员工编号",
    "姓名",
    "企业邮箱",
    "一级部门",
    "二级部门",
    "员工类别编码",
    "企业微信UserId",
)
RELATIONSHIP_HEADERS = ("被评价人工号", "评价人工号", "关系类型")
XLSX_ENCRYPTION_FLAGS = 0x1 | 0x40 | 0x2000
XLSX_LOCAL_HEADER = struct.Struct("<4s5H3I2H")


class ImportPreviewError(ValueError):
    pass


class ImportCommitError(ValueError):
    pass


@dataclass(frozen=True)
class ImportCommitResult:
    created_count: int = 0
    updated_count: int = 0
    skipped_duplicate_count: int = 0
    deactivated_count: int = 0


def _safe_filename(filename: str) -> str:
    return Path(str(filename).replace("\\", "/")).name[:255]


def validate_upload(content: bytes, filename: str) -> str:
    extension = Path(filename).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise ImportPreviewError("仅支持 .xlsx 或 .csv 文件")
    if len(content) > MAX_UPLOAD_BYTES:
        raise ImportPreviewError("上传文件不能超过 10MB")
    if not content:
        raise ImportPreviewError("上传文件不能为空")
    return extension


def _is_safe_xlsx_member_name(name: str) -> bool:
    if not name or "\x00" in name or "\\" in name or name.startswith("/"):
        return False
    path = name[:-1] if name.endswith("/") else name
    if not path:
        return False
    parts = path.split("/")
    if any(part in {"", ".", ".."} for part in parts):
        return False
    return not (len(parts[0]) >= 2 and parts[0][1] == ":")


def _xlsx_member_data_bounds(content: bytes, member):
    offset = member.header_offset
    if offset < 0 or offset + XLSX_LOCAL_HEADER.size > len(content):
        raise ImportPreviewError("文件无法解析")
    (
        signature,
        _version,
        flags,
        compression,
        _modified_time,
        _modified_date,
        _crc,
        _compressed_size,
        _uncompressed_size,
        filename_length,
        extra_length,
    ) = XLSX_LOCAL_HEADER.unpack_from(content, offset)
    if signature != b"PK\x03\x04":
        raise ImportPreviewError("文件无法解析")
    filename_start = offset + XLSX_LOCAL_HEADER.size
    filename_end = filename_start + filename_length
    data_start = filename_end + extra_length
    data_end = data_start + member.compress_size
    if (
        filename_end > len(content)
        or data_start > len(content)
        or data_end > len(content)
    ):
        raise ImportPreviewError("文件无法解析")
    encoding = "utf-8" if flags & 0x800 else "cp437"
    local_name = content[filename_start:filename_end].decode(encoding)
    if local_name != member.filename or compression != member.compress_type:
        raise ImportPreviewError("文件无法解析")
    if (flags | member.flag_bits) & XLSX_ENCRYPTION_FLAGS:
        raise ImportPreviewError("XLSX 不能包含加密文件项")
    return data_start, data_end


def _copy_stored_xlsx_member(content, start, end, destination, expanded_so_far):
    actual_size = 0
    actual_crc = 0
    for offset in range(start, end, XLSX_STREAM_CHUNK_BYTES):
        chunk = memoryview(content)[offset : min(offset + XLSX_STREAM_CHUNK_BYTES, end)]
        actual_size += len(chunk)
        if expanded_so_far + actual_size > MAX_XLSX_UNCOMPRESSED_BYTES:
            raise ImportPreviewError("XLSX 解压后不能超过 20MB")
        actual_crc = zlib.crc32(chunk, actual_crc)
        destination.write(chunk)
    return actual_size, actual_crc


def _copy_deflated_xlsx_member(content, start, end, destination, expanded_so_far):
    decompressor = zlib.decompressobj(-zlib.MAX_WBITS)
    actual_size = 0
    actual_crc = 0
    cursor = start
    pending = b""
    while True:
        if not pending and cursor < end:
            next_cursor = min(cursor + XLSX_STREAM_CHUNK_BYTES, end)
            pending = memoryview(content)[cursor:next_cursor]
            cursor = next_cursor
        remaining_limit = MAX_XLSX_UNCOMPRESSED_BYTES - expanded_so_far - actual_size
        output_limit = min(XLSX_STREAM_CHUNK_BYTES, remaining_limit + 1)
        output = decompressor.decompress(pending, output_limit)
        pending = decompressor.unconsumed_tail
        if output:
            actual_size += len(output)
            if expanded_so_far + actual_size > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise ImportPreviewError("XLSX 解压后不能超过 20MB")
            actual_crc = zlib.crc32(output, actual_crc)
            destination.write(output)
        if decompressor.eof:
            if decompressor.unused_data or pending or cursor != end:
                raise ImportPreviewError("文件无法解析")
            break
        if not output and not pending and cursor == end:
            raise ImportPreviewError("文件无法解析")
    return actual_size, actual_crc


def _sanitized_member_info(member):
    sanitized = ZipInfo(member.filename, date_time=member.date_time)
    sanitized.compress_type = ZIP_STORED
    sanitized.external_attr = member.external_attr
    sanitized.create_system = member.create_system
    return sanitized


def _validate_and_sanitize_xlsx_archive(content: bytes) -> BytesIO:
    output = BytesIO()
    try:
        with ZipFile(BytesIO(content)) as source:
            members = source.infolist()
            if len(members) > MAX_XLSX_MEMBERS:
                raise ImportPreviewError("XLSX 压缩包文件项过多")
            seen_names = set()
            expanded_bytes = 0
            with ZipFile(output, "w", compression=ZIP_STORED) as sanitized:
                for member in members:
                    if member.filename in seen_names:
                        raise ImportPreviewError("XLSX 压缩包包含重复文件项")
                    seen_names.add(member.filename)
                    if not _is_safe_xlsx_member_name(member.filename):
                        raise ImportPreviewError("XLSX 压缩包文件项路径不安全")
                    if member.flag_bits & XLSX_ENCRYPTION_FLAGS:
                        raise ImportPreviewError("XLSX 不能包含加密文件项")
                    if member.compress_type not in {ZIP_STORED, ZIP_DEFLATED}:
                        raise ImportPreviewError("XLSX 包含不支持的压缩方式")
                    start, end = _xlsx_member_data_bounds(content, member)
                    with sanitized.open(_sanitized_member_info(member), "w") as target:
                        if member.compress_type == ZIP_STORED:
                            actual_size, actual_crc = _copy_stored_xlsx_member(
                                content, start, end, target, expanded_bytes
                            )
                        else:
                            actual_size, actual_crc = _copy_deflated_xlsx_member(
                                content, start, end, target, expanded_bytes
                            )
                    expanded_bytes += actual_size
                    if actual_size != member.file_size or actual_crc != member.CRC:
                        raise ImportPreviewError("文件无法解析")
        output.seek(0)
        return output
    except ImportPreviewError:
        output.close()
        raise
    except (BadZipFile, OSError, UnicodeDecodeError, ValueError, zlib.error) as exc:
        output.close()
        raise ImportPreviewError("文件无法解析") from exc


def _read_rows(content: bytes, filename: str, required_headers: tuple[str, ...]):
    extension = validate_upload(content, filename)
    workbook = None
    sanitized_archive = None
    try:
        if extension == ".csv":
            raw_rows = csv.reader(StringIO(content.decode("utf-8-sig")))
        else:
            sanitized_archive = _validate_and_sanitize_xlsx_archive(content)
            workbook = load_workbook(
                sanitized_archive, read_only=True, data_only=True
            )
            worksheet = workbook.active
            raw_rows = worksheet.iter_rows(values_only=True)
        headers = tuple(str(value or "").strip() for value in next(raw_rows, ()))
        if headers != required_headers or len(set(headers)) != len(headers):
            raise ImportPreviewError("导入列必须完全匹配固定列且不得重复")

        for data_row_number, raw_values in enumerate(raw_rows, start=1):
            if data_row_number > MAX_IMPORT_ROWS:
                raise ImportPreviewError("导入数据最多 500 行")
            if len(raw_values) != len(headers):
                raise ImportPreviewError("导入数据每行必须严格包含固定列")
            normalized = {
                header: str(value or "").strip()
                for header, value in zip(headers, raw_values, strict=False)
            }
            if any(normalized.values()):
                yield data_row_number + 1, normalized
    except ImportPreviewError:
        raise
    except (UnicodeDecodeError, csv.Error, OSError, ValueError, KeyError) as exc:
        raise ImportPreviewError("文件无法解析") from exc
    except Exception as exc:
        raise ImportPreviewError("文件无法解析") from exc
    finally:
        if workbook is not None:
            workbook.close()
        if sanitized_archive is not None:
            sanitized_archive.close()


def _issue(row_number: int, error: RosterValidationError | str, code: str | None = None):
    if isinstance(error, RosterValidationError):
        return ImportIssue(
            row_number=row_number,
            code=error.code,
            message=f"第 {row_number} 行：{error}",
        )
    return ImportIssue(
        row_number=row_number,
        code=code or "INVALID",
        message=f"第 {row_number} 行：{error}",
    )


def _append_issue(issues, issue):
    issues.append(issue)
    if len(issues) > MAX_IMPORT_ISSUES:
        raise ImportPreviewError("导入异常行数不能超过 100")


def _fingerprint(values) -> str:
    payload = json.dumps(values, ensure_ascii=False, separators=(",", ":"), sort_keys=True)
    return sha256(payload.encode("utf-8")).hexdigest()


def _employee_fingerprint(employee: Employee) -> str:
    return _fingerprint(
        {
            "id": employee.pk,
            "employee_no": employee.employee_no,
            "name": employee.name,
            "corporate_email": employee.corporate_email,
            "department_level_1": employee.department_level_1,
            "department_level_2": employee.department_level_2,
            "category_id": employee.category_id,
            "manager_id": employee.manager_id,
            "wecom_userid": employee.wecom_userid,
            "user_id": employee.user_id,
            "is_active": employee.is_active,
        }
    )


def _category_fingerprint(category: EmployeeCategory) -> str:
    return _fingerprint(
        {
            "id": category.pk,
            "code": category.code,
            "name": category.name,
            "is_active": category.is_active,
        }
    )


def _relationship_fingerprint(relationship: EvaluationRelationship) -> str:
    return _fingerprint(
        {
            "id": relationship.pk,
            "subject_id": relationship.subject_id,
            "evaluator_id": relationship.evaluator_id,
            "relationship_type": relationship.relationship_type,
            "is_active": relationship.is_active,
        }
    )


def _save_preview_batch(
    *, import_type, content, filename, actor, rows, issues, duplicate_count, preview
):
    batch = ImportBatch.objects.create(
        import_type=import_type,
        original_filename=_safe_filename(filename),
        file_sha256=sha256(content).hexdigest(),
        rows=rows,
        preview=preview,
        valid_count=len(rows),
        duplicate_count=duplicate_count,
        issue_count=len(issues),
        create_count=len(preview.get("create_items", [])),
        update_count=len(preview.get("update_items", [])),
        skip_count=len(preview.get("skip_items", [])),
        deactivate_count=len(
            preview.get("deactivate_employee_nos", [])
            or preview.get("deactivate_relationship_ids", [])
        ),
        created_by=actor,
    )
    for issue in issues:
        issue.batch = batch
    ImportIssue.objects.bulk_create(issues)
    return batch


@transaction.atomic
def preview_roster_upload(content: bytes, filename: str, actor) -> ImportBatch:
    require_hr_actor(actor)
    raw_rows = list(_read_rows(content, filename, ROSTER_HEADERS))
    rows = []
    issues = []
    duplicate_count = 0
    seen_employee_nos = set()
    seen_emails = set()
    seen_wecom_userids = set()

    category_codes = {row["员工类别编码"] for _, row in raw_rows}
    categories = {
        category.code: category
        for category in EmployeeCategory.objects.filter(code__in=category_codes)
    }
    requested_employee_nos = {row["员工编号"] for _, row in raw_rows}
    requested_emails = {row["企业邮箱"].lower() for _, row in raw_rows}
    requested_wecom_userids = {
        row["企业微信UserId"] for _, row in raw_rows if row["企业微信UserId"]
    }
    identity_employees = list(
        Employee.objects.filter(
            Q(employee_no__in=requested_employee_nos)
            | Q(corporate_email__in=requested_emails)
            | Q(wecom_userid__in=requested_wecom_userids)
        )
    )
    active_employees = list(Employee.objects.filter(is_active=True).order_by("employee_no"))
    employees_by_no = {employee.employee_no: employee for employee in identity_employees}
    employees_by_email = {
        employee.corporate_email.lower(): employee for employee in identity_employees
    }
    employees_by_wecom = {
        employee.wecom_userid: employee
        for employee in identity_employees
        if employee.wecom_userid
    }

    for row_number, row in raw_rows:
        category = categories.get(row["员工类别编码"])
        try:
            values = normalize_employee_fields(
                employee_no=row["员工编号"],
                name=row["姓名"],
                corporate_email=row["企业邮箱"],
                department_level_1=row["一级部门"],
                department_level_2=row["二级部门"],
                wecom_userid=row["企业微信UserId"],
            )
            if category is None:
                raise RosterValidationError("员工类别不存在", "CATEGORY_NOT_FOUND")
            if not category.is_active:
                raise RosterValidationError("员工类别已停用", "CATEGORY_INACTIVE")
        except RosterValidationError as exc:
            _append_issue(issues, _issue(row_number, exc))
            continue

        employee_no = values["employee_no"]
        email = values["corporate_email"]
        wecom_userid = values["wecom_userid"]
        existing = employees_by_no.get(employee_no)
        if employee_no in seen_employee_nos:
            _append_issue(
                issues,
                _issue(row_number, "员工编号在文件内重复", "DUPLICATE_EMPLOYEE_NO"),
            )
            continue
        if email in seen_emails:
            _append_issue(
                issues,
                _issue(row_number, "企业邮箱在文件内重复", "DUPLICATE_CORPORATE_EMAIL"),
            )
            continue
        if wecom_userid and wecom_userid in seen_wecom_userids:
            _append_issue(
                issues,
                _issue(row_number, "企业微信 UserId 在文件内重复", "DUPLICATE_WECOM_USERID"),
            )
            continue
        email_conflict = employees_by_email.get(email)
        if email_conflict and (existing is None or email_conflict.pk != existing.pk):
            _append_issue(
                issues,
                _issue(row_number, "企业邮箱已存在", "DUPLICATE_CORPORATE_EMAIL"),
            )
            continue
        wecom_conflict = employees_by_wecom.get(wecom_userid) if wecom_userid else None
        if wecom_conflict and (existing is None or wecom_conflict.pk != existing.pk):
            _append_issue(
                issues,
                _issue(row_number, "企业微信 UserId 已存在", "DUPLICATE_WECOM_USERID"),
            )
            continue

        seen_employee_nos.add(values["employee_no"])
        seen_emails.add(values["corporate_email"])
        if values["wecom_userid"]:
            seen_wecom_userids.add(values["wecom_userid"])
        if existing:
            duplicate_count += 1
        rows.append(
            {
                "row_number": row_number,
                "employee_no": values["employee_no"],
                "name": values["name"],
                "corporate_email": values["corporate_email"],
                "department_level_1": values["department_level_1"],
                "department_level_2": values["department_level_2"],
                "category_id": category.pk,
                "wecom_userid": values["wecom_userid"],
                "existing_employee_id": existing.pk if existing else None,
                "existing_employee_fingerprint": (
                    _employee_fingerprint(existing) if existing else None
                ),
                "category_fingerprint": _category_fingerprint(category),
            }
        )

    incoming_nos = {row["employee_no"] for row in rows}
    deactivation_employees = [
        {
            "id": employee.pk,
            "employee_no": employee.employee_no,
            "name": employee.name,
            "department_level_1": employee.department_level_1,
            "department_level_2": employee.department_level_2,
            "fingerprint": _employee_fingerprint(employee),
        }
        for employee in active_employees
        if employee.employee_no not in incoming_nos
    ]
    preview = {
        "create_items": [row["employee_no"] for row in rows if not row["existing_employee_id"]],
        "update_items": [row["employee_no"] for row in rows if row["existing_employee_id"]],
        "skip_items": [row["employee_no"] for row in rows if row["existing_employee_id"]],
        "deactivate_employee_nos": [
            employee["employee_no"] for employee in deactivation_employees
        ],
        "deactivate_employee_ids": [
            employee["id"] for employee in deactivation_employees
        ],
        "deactivate_employees": deactivation_employees,
        "deactivation_display_fingerprint": _fingerprint(deactivation_employees),
    }
    return _save_preview_batch(
        import_type=ImportBatch.Type.ROSTER,
        content=content,
        filename=filename,
        actor=actor,
        rows=rows,
        issues=issues,
        duplicate_count=duplicate_count,
        preview=preview,
    )


@transaction.atomic
def preview_relationship_upload(content: bytes, filename: str, actor) -> ImportBatch:
    require_hr_actor(actor)
    raw_rows = list(_read_rows(content, filename, RELATIONSHIP_HEADERS))
    rows = []
    issues = []
    duplicate_count = 0
    seen_keys = set()

    requested_employee_nos = {
        employee_no
        for _, row in raw_rows
        for employee_no in (row["被评价人工号"], row["评价人工号"])
    }
    employees = {
        employee.employee_no: employee
        for employee in Employee.objects.filter(
            employee_no__in=requested_employee_nos, is_active=True
        )
    }
    relationships = list(
        EvaluationRelationship.objects.select_related("subject", "evaluator").order_by("id")
    )
    relationships_by_key = {}
    for relationship in relationships:
        key = (
            relationship.subject_id,
            relationship.evaluator_id,
            relationship.relationship_type,
        )
        relationships_by_key.setdefault(key, []).append(relationship)

    for row_number, row in raw_rows:
        relationship_type = row["关系类型"]
        subject = employees.get(row["被评价人工号"])
        evaluator = employees.get(row["评价人工号"])
        try:
            if subject is None or evaluator is None:
                raise RosterValidationError(
                    "员工编号不存在或已停用", "EMPLOYEE_NOT_FOUND"
                )
            validate_relationship_entities(subject, evaluator, relationship_type)
        except RosterValidationError as exc:
            _append_issue(issues, _issue(row_number, exc))
            continue
        key = (subject.pk, evaluator.pk, relationship_type)
        if key in seen_keys:
            _append_issue(
                issues,
                _issue(row_number, "协作关系在文件内重复", "DUPLICATE_RELATION"),
            )
            continue
        seen_keys.add(key)
        matching = relationships_by_key.get(key, [])
        active = next((item for item in matching if item.is_active), None)
        inactive = next((item for item in matching if not item.is_active), None)
        existing = active or inactive
        if active:
            duplicate_count += 1
        rows.append(
            {
                "row_number": row_number,
                "subject_no": subject.employee_no,
                "subject_id": subject.pk,
                "subject_fingerprint": _employee_fingerprint(subject),
                "evaluator_no": evaluator.employee_no,
                "evaluator_id": evaluator.pk,
                "evaluator_fingerprint": _employee_fingerprint(evaluator),
                "relationship_type": relationship_type,
                "existing_relationship_id": existing.pk if existing else None,
                "existing_active": bool(active),
                "existing_relationship_fingerprint": (
                    _relationship_fingerprint(existing) if existing else None
                ),
            }
        )

    incoming_keys = {
        (row["subject_no"], row["evaluator_no"], row["relationship_type"])
        for row in rows
    }
    deactivation_relationships = [
        {
            "id": relation.pk,
            "subject_id": relation.subject_id,
            "subject_no": relation.subject.employee_no,
            "subject_name": relation.subject.name,
            "subject_fingerprint": _employee_fingerprint(relation.subject),
            "evaluator_id": relation.evaluator_id,
            "evaluator_no": relation.evaluator.employee_no,
            "evaluator_name": relation.evaluator.name,
            "evaluator_fingerprint": _employee_fingerprint(relation.evaluator),
            "relationship_type": relation.relationship_type,
            "relationship_type_label": relation.get_relationship_type_display(),
            "fingerprint": _relationship_fingerprint(relation),
        }
        for relation in EvaluationRelationship.objects.filter(is_active=True)
        .select_related("subject", "evaluator")
        .order_by("id")
        if (
            relation.subject.employee_no,
            relation.evaluator.employee_no,
            relation.relationship_type,
        )
        not in incoming_keys
    ]
    preview = {
        "create_items": [row["row_number"] for row in rows if not row["existing_relationship_id"]],
        "update_items": [row["row_number"] for row in rows if row["existing_relationship_id"]],
        "skip_items": [row["row_number"] for row in rows if row["existing_active"]],
        "deactivate_relationship_ids": [
            relationship["id"] for relationship in deactivation_relationships
        ],
        "deactivate_relationships": deactivation_relationships,
        "deactivation_display_fingerprint": _fingerprint(deactivation_relationships),
    }
    return _save_preview_batch(
        import_type=ImportBatch.Type.RELATIONSHIP,
        content=content,
        filename=filename,
        actor=actor,
        rows=rows,
        issues=issues,
        duplicate_count=duplicate_count,
        preview=preview,
    )


def _validate_commit_options(batch, mode, duplicate_policy):
    if mode not in {"replace", "append"}:
        raise ImportCommitError("导入模式必须为 replace 或 append")
    if duplicate_policy not in {"update", "skip"}:
        if batch.duplicate_count:
            raise ImportCommitError("存在重复记录时必须明确选择重复处理策略")
        raise ImportCommitError("重复处理策略必须为 update 或 skip")
    if batch.status != ImportBatch.Status.PREVIEWED:
        raise ImportCommitError("导入批次已提交")
    if batch.issue_count:
        raise ImportCommitError("导入批次存在异常，不能提交")


def has_complete_deactivation_details(batch) -> bool:
    if batch.deactivate_count == 0:
        return True
    if batch.import_type == ImportBatch.Type.ROSTER:
        details_key = "deactivate_employees"
        ids_key = "deactivate_employee_ids"
        required_fields = {
            "id",
            "employee_no",
            "name",
            "department_level_1",
            "department_level_2",
            "fingerprint",
        }
    else:
        details_key = "deactivate_relationships"
        ids_key = "deactivate_relationship_ids"
        required_fields = {
            "id",
            "subject_id",
            "subject_no",
            "subject_name",
            "subject_fingerprint",
            "evaluator_id",
            "evaluator_no",
            "evaluator_name",
            "evaluator_fingerprint",
            "relationship_type",
            "relationship_type_label",
            "fingerprint",
        }
    details = batch.preview.get(details_key)
    expected_ids = batch.preview.get(ids_key)
    expected_display_fingerprint = batch.preview.get(
        "deactivation_display_fingerprint"
    )
    return not (
        not isinstance(details, list)
        or not isinstance(expected_ids, list)
        or not isinstance(expected_display_fingerprint, str)
        or len(details) != batch.deactivate_count
        or any(not isinstance(item, dict) for item in details)
        or [item.get("id") for item in details] != expected_ids
        or any(not required_fields.issubset(item) for item in details)
        or _fingerprint(details) != expected_display_fingerprint
    )


def _complete_deactivation_details(batch, details_key, ids_key):
    if not has_complete_deactivation_details(batch):
        raise ImportCommitError("预览缺少完整停用明细，请重新预检")
    return batch.preview.get(details_key) or []


def _preflight_roster(batch, mode):
    # Shared mutable dependency order is Employee -> EmployeeCategory ->
    # EvaluationRelationship; see apps.roster.locking.
    locked_employees = list(
        Employee.objects.select_for_update(of=("self",)).order_by("pk")
    )
    employees_by_id = {employee.pk: employee for employee in locked_employees}
    employees_by_no = {employee.employee_no: employee for employee in locked_employees}
    employees_by_email = {
        employee.corporate_email.lower(): employee for employee in locked_employees
    }
    employees_by_wecom = {
        employee.wecom_userid: employee
        for employee in locked_employees
        if employee.wecom_userid
    }
    categories = {
        category.pk: category
        for category in EmployeeCategory.objects.select_for_update(of=("self",))
        .filter(pk__in={row["category_id"] for row in batch.rows})
        .order_by("pk")
    }

    for row in batch.rows:
        category = categories.get(row["category_id"])
        if (
            category is None
            or row.get("category_fingerprint") != _category_fingerprint(category)
        ):
            raise ImportCommitError("业务数据已变化，请重新预检")
        existing = employees_by_no.get(row["employee_no"])
        preview_id = row["existing_employee_id"]
        if (existing.pk if existing else None) != preview_id:
            raise ImportCommitError("业务数据已变化，请重新预检")
        if existing:
            if row.get("existing_employee_fingerprint") != _employee_fingerprint(existing):
                raise ImportCommitError("业务数据已变化，请重新预检")
            continue
        if employees_by_email.get(row["corporate_email"].lower()) is not None:
            raise ImportCommitError("业务数据已变化，请重新预检")
        wecom_userid = row.get("wecom_userid")
        if wecom_userid and employees_by_wecom.get(wecom_userid) is not None:
            raise ImportCommitError("业务数据已变化，请重新预检")

    deactivation_details = []
    if mode == "replace":
        deactivation_details = _complete_deactivation_details(
            batch, "deactivate_employees", "deactivate_employee_ids"
        )
        for detail in deactivation_details:
            employee = employees_by_id.get(detail["id"])
            if (
                employee is None
                or employee.employee_no != detail["employee_no"]
                or not employee.is_active
                or detail.get("fingerprint") != _employee_fingerprint(employee)
            ):
                raise ImportCommitError("业务数据已变化，请重新预检")
    return employees_by_no, deactivation_details


def _roster_commit(batch, mode, duplicate_policy, actor):
    employees, deactivation_details = _preflight_roster(batch, mode)
    created = updated = skipped = 0
    for row in batch.rows:
        existing = employees.get(row["employee_no"])
        fields = {
            key: row[key]
            for key in (
                "employee_no",
                "name",
                "corporate_email",
                "department_level_1",
                "department_level_2",
                "category_id",
                "wecom_userid",
            )
        }
        try:
            if existing:
                if duplicate_policy == "skip":
                    skipped += 1
                    continue
                update_employee_record(existing.pk, actor=actor, **fields)
                updated += 1
            else:
                created_employee = create_employee_record(actor=actor, **fields)
                employees[created_employee.employee_no] = created_employee
                created += 1
        except (RosterValidationError, IntegrityError) as exc:
            raise ImportCommitError("业务数据已变化，请重新预检") from exc

    deactivated = 0
    if mode == "replace":
        deactivated = Employee.objects.filter(
            pk__in=[detail["id"] for detail in deactivation_details], is_active=True
        ).update(is_active=False)
    return ImportCommitResult(created, updated, skipped, deactivated)


def _preflight_relationships(batch, mode):
    employees = {
        employee.pk: employee
        for employee in Employee.objects.select_for_update(of=("self",)).order_by(
            "pk"
        )
    }
    locked_relationships = list(
        EvaluationRelationship.objects.select_for_update(of=("self",)).order_by("pk")
    )
    relationships = {relation.pk: relation for relation in locked_relationships}
    relationships_by_key = {}
    for relationship in locked_relationships:
        key = (
            relationship.subject_id,
            relationship.evaluator_id,
            relationship.relationship_type,
        )
        relationships_by_key.setdefault(key, []).append(relationship)

    for row in batch.rows:
        subject = employees.get(row.get("subject_id"))
        evaluator = employees.get(row.get("evaluator_id"))
        if (
            subject is None
            or evaluator is None
            or subject.employee_no != row["subject_no"]
            or evaluator.employee_no != row["evaluator_no"]
            or row.get("subject_fingerprint") != _employee_fingerprint(subject)
            or row.get("evaluator_fingerprint") != _employee_fingerprint(evaluator)
        ):
            raise ImportCommitError("业务数据已变化，请重新预检")
        key = (subject.pk, evaluator.pk, row["relationship_type"])
        matching = relationships_by_key.get(key, [])
        preview_id = row["existing_relationship_id"]
        if preview_id is None:
            if matching:
                raise ImportCommitError("业务数据已变化，请重新预检")
            continue
        existing = relationships.get(preview_id)
        if (
            existing is None
            or existing not in matching
            or row.get("existing_relationship_fingerprint")
            != _relationship_fingerprint(existing)
            or bool(existing.is_active) != bool(row["existing_active"])
        ):
            raise ImportCommitError("业务数据已变化，请重新预检")
        if not row["existing_active"] and any(item.is_active for item in matching):
            raise ImportCommitError("业务数据已变化，请重新预检")

    deactivation_details = []
    if mode == "replace":
        deactivation_details = _complete_deactivation_details(
            batch, "deactivate_relationships", "deactivate_relationship_ids"
        )
        for detail in deactivation_details:
            relationship = relationships.get(detail["id"])
            subject = employees.get(detail.get("subject_id"))
            evaluator = employees.get(detail.get("evaluator_id"))
            if (
                relationship is None
                or not relationship.is_active
                or detail.get("fingerprint") != _relationship_fingerprint(relationship)
                or subject is None
                or evaluator is None
                or detail.get("subject_fingerprint") != _employee_fingerprint(subject)
                or detail.get("evaluator_fingerprint") != _employee_fingerprint(evaluator)
            ):
                raise ImportCommitError("业务数据已变化，请重新预检")
    return relationships, deactivation_details


def _relationship_commit(batch, mode, duplicate_policy, actor):
    relationships, deactivation_details = _preflight_relationships(batch, mode)
    created = updated = skipped = 0
    for row in batch.rows:
        preview_id = row["existing_relationship_id"]
        existing = relationships.get(preview_id) if preview_id else None
        if row["existing_active"]:
            if duplicate_policy == "skip":
                skipped += 1
            else:
                updated += 1
            continue
        try:
            relation = upsert_relationship(
                subject_no=row["subject_no"],
                evaluator_no=row["evaluator_no"],
                relationship_type=row["relationship_type"],
                actor=actor,
            )
        except (RosterValidationError, IntegrityError) as exc:
            raise ImportCommitError("业务数据已变化，请重新预检") from exc
        relationships[relation.pk] = relation
        if preview_id:
            updated += 1
        else:
            created += 1

    deactivated = 0
    if mode == "replace":
        deactivated = EvaluationRelationship.objects.filter(
            pk__in=[detail["id"] for detail in deactivation_details], is_active=True
        ).update(is_active=False)
    return ImportCommitResult(created, updated, skipped, deactivated)


@transaction.atomic
def commit_import_batch(batch_id, mode, duplicate_policy, actor) -> ImportCommitResult:
    require_hr_actor(actor)
    try:
        batch = ImportBatch.objects.select_for_update(of=("self",)).get(
            public_id=batch_id
        )
    except ImportBatch.DoesNotExist as exc:
        raise ImportCommitError("导入批次不存在") from exc
    _validate_commit_options(batch, mode, duplicate_policy)

    if batch.import_type == ImportBatch.Type.ROSTER:
        result = _roster_commit(batch, mode, duplicate_policy, actor)
    else:
        result = _relationship_commit(batch, mode, duplicate_policy, actor)
    batch.status = ImportBatch.Status.COMMITTED
    batch.committed_by = actor
    batch.committed_at = timezone.now()
    batch.commit_mode = mode
    batch.duplicate_policy = duplicate_policy
    batch.save(
        update_fields=[
            "status",
            "committed_by",
            "committed_at",
            "commit_mode",
            "duplicate_policy",
        ]
    )
    action = (
        "ROSTER_IMPORT_COMMITTED"
        if batch.import_type == ImportBatch.Type.ROSTER
        else "RELATIONSHIP_IMPORT_COMMITTED"
    )
    record_audit(
        actor,
        action,
        batch,
        {
            "created_count": result.created_count,
            "updated_count": result.updated_count,
            "skipped_count": result.skipped_duplicate_count,
            "deactivated_count": result.deactivated_count,
            "mode": mode,
            "duplicate_policy": duplicate_policy,
        },
        idempotency_key=f"import:{batch.public_id}:committed",
    )
    return result
